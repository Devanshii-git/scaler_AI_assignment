"""
Integration test for full PiiRedactionPipeline.
"""

import os
import docx
import pytest
from src.pipeline import PiiRedactionPipeline


def test_pipeline_end_to_end(tmp_path):
    input_path = str(tmp_path / "prospectus_snippet.docx")
    output_path = str(tmp_path / "prospectus_snippet_redacted.docx")

    doc = docx.Document()
    # Paragraph with Promoter details
    p1 = doc.add_paragraph("BOARD OF DIRECTORS:")
    p2 = doc.add_paragraph("Mr. Rajesh Kushal Hegde is the Promoter and Managing Director of KSH International Limited.")
    p3 = doc.add_paragraph("Registered Office: Plot No. 12, Chakan Industrial Area, MIDC, Pune 410501, Maharashtra, India. Tel: +91 9876543210.")
    p4 = doc.add_paragraph("For queries contact Sarthak Malvadkar at compliance@kshgroup.com.")
    p5 = doc.add_paragraph("Director Date of Birth: 15 May 1975. This Draft Prospectus is Dated December 10, 2025.")

    doc.save(input_path)

    # Run pipeline
    pipeline = PiiRedactionPipeline(enable_neural=True, model_path="models/gliner_multi_pii")
    summary = pipeline.process_document(input_path, output_path)

    assert os.path.exists(output_path)
    assert summary["total_pii_detected"] >= 4

    types_detected = summary["pii_by_type"]
    assert "EMAIL" in types_detected
    assert "PHONE" in types_detected

    # Inspect redacted document text
    red_doc = docx.Document(output_path)
    full_text = " ".join([p.text for p in red_doc.paragraphs])

    # Assert raw PII strings are absent
    assert "compliance@kshgroup.com" not in full_text
    assert "+91 9876543210" not in full_text
    assert "Rajesh Kushal Hegde" not in full_text
