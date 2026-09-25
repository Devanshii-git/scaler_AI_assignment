"""
Unit test for DocxReconstructor.
"""

import os
import docx
import pytest
from src.document.docx_reconstructor import DocxReconstructor
from src.pseudonymization.synthesizer import PseudonymSynthesizer
from src.ingestion.docx_parser import DocxParser
from src.detectors.base import EntitySpan


def test_docx_reconstruction_formatting_preservation(tmp_path):
    # Create input test DOCX
    input_path = str(tmp_path / "test_input.docx")
    output_path = str(tmp_path / "test_output.docx")

    doc = docx.Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Contact ")
    r2 = p.add_run("Rajesh Hegde")
    r2.bold = True
    r3 = p.add_run(" at ")
    r4 = p.add_run("compliance@kshgroup.com")
    r4.italic = True
    doc.save(input_path)

    # Ingest document
    parser = DocxParser()
    blocks, images, loaded_doc = parser.parse_document(input_path)
    assert len(blocks) >= 1

    synth = PseudonymSynthesizer()
    reconstructor = DocxReconstructor(synthesizer=synth)

    # Fake detected spans for block
    block = blocks[0]
    name_span = EntitySpan(
        entity_id="e1",
        type="PERSON",
        text="Rajesh Hegde",
        start=block.normalized_text.index("Rajesh Hegde"),
        end=block.normalized_text.index("Rajesh Hegde") + len("Rajesh Hegde")
    )
    email_span = EntitySpan(
        entity_id="e2",
        type="EMAIL",
        text="compliance@kshgroup.com",
        start=block.normalized_text.index("compliance@kshgroup.com"),
        end=block.normalized_text.index("compliance@kshgroup.com") + len("compliance@kshgroup.com")
    )

    block_spans = {block.block_id: [name_span, email_span]}

    reconstructor.redact_document(
        original_docx_path=input_path,
        output_docx_path=output_path,
        blocks=blocks,
        block_spans=block_spans,
        doc=loaded_doc
    )

    assert os.path.exists(output_path)

    # Inspect resulting document
    res_doc = docx.Document(output_path)
    res_p = res_doc.paragraphs[0]

    # Verify PII is gone
    assert "Rajesh Hegde" not in res_p.text
    assert "compliance@kshgroup.com" not in res_p.text

    # Verify styling exists
    bold_runs = [r for r in res_p.runs if r.bold and r.text.strip()]
    italic_runs = [r for r in res_p.runs if r.italic and r.text.strip()]

    assert len(bold_runs) >= 1
    assert len(italic_runs) >= 1
